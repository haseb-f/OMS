"""
DOCX rendering of the manual with python-docx — real right-to-left document:
section <w:bidi/>, every paragraph <w:bidi/> (alignment = paragraph start, i.e. right), Arabic runs
<w:rtl/> with Arial as the complex-script font, tables <w:bidiVisual/>, Word heading styles with
bookmarks, a clickable table of contents (internal hyperlinks + PAGEREF fields), header/footer
with the release, SHA and PAGE / NUMPAGES fields, and a cover page with the logo.
"""
import struct
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from meta import META, ROOT, figure_width_mm

NAVY = "0F2747"
BLUE = "3B6FD8"
TEAL = "3BB5C4"
ORANGE = "E8590C"
GREY = "5B6B82"
FONT = "Arial"
CODE_FONT = "Consolas"
SHOTS = ROOT / "screenshots" / "annotated"
ASSETS = ROOT / "assets"
BOX_FILL = {"note": "EEF4FD", "tip": "EAF7EF", "warn": "FFF5E6"}
BOX_LINE = {"note": BLUE, "tip": "2E9D5B", "warn": "E08A00"}

_bookmark_id = [100]


# ───────────── low-level helpers ─────────────
def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


def set_bidi(paragraph, center=False):
    pPr = paragraph._p.get_or_add_pPr()
    if pPr.find(qn("w:bidi")) is None:
        pPr.insert(0, _el("w:bidi"))
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return paragraph


def style_run(run, size=None, bold=None, color=None, rtl=True, font=FONT, italic=None):
    rPr = run._r.get_or_add_rPr()
    fonts = rPr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = _el("w:rFonts")
        rPr.insert(0, fonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(a), font)
    if bold is not None:
        run.bold = bold
        if bold:
            rPr.append(_el("w:bCs"))
    if italic:
        run.italic = True
    if size:
        run.font.size = Pt(size)
        rPr.append(_el("w:szCs", **{"w:val": int(size * 2)}))
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if rtl:
        rPr.append(_el("w:rtl"))
    lang = _el("w:lang", **{"w:val": "en-US", "w:bidi": "ar-EG"})
    rPr.append(lang)
    return run


def add_runs(paragraph, runs, size=10.5, color=None, bold=False):
    for r in runs:
        if r.get("code"):
            run = paragraph.add_run(r["t"])
            style_run(run, size=size - 1, bold=r.get("b") or bold, color=color or "1B2533", rtl=False, font=CODE_FONT)
            shd = _el("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": "EEF2F8"})
            run._r.get_or_add_rPr().append(shd)
        else:
            run = paragraph.add_run(r["t"])
            style_run(run, size=size, bold=r.get("b") or bold, color=(NAVY if r.get("b") and not color else color))
    return paragraph


def para(doc_or_cell, runs=None, size=10.5, space_after=4, keep_next=False, center=False, color=None, bold=False):
    p = doc_or_cell.add_paragraph()
    set_bidi(p, center)
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(0)
    pf.line_spacing = 1.25
    if keep_next:
        pf.keep_with_next = True
    if runs:
        add_runs(p, runs, size=size, color=color, bold=bold)
    return p


def indent_start(p, start_twips, hanging_twips=0):
    pPr = p._p.get_or_add_pPr()
    ind = _el("w:ind", **{"w:start": start_twips, "w:hanging": hanging_twips})
    pPr.append(ind)


def add_bookmark(paragraph, name):
    bid = _bookmark_id[0]
    _bookmark_id[0] += 1
    start = _el("w:bookmarkStart", **{"w:id": bid, "w:name": name})
    end = _el("w:bookmarkEnd", **{"w:id": bid})
    paragraph._p.insert(1, start)
    paragraph._p.append(end)


def add_field(paragraph, instr, cached="", size=9, color=GREY):
    def r(child):
        run = paragraph.add_run()
        style_run(run, size=size, color=color, rtl=False)
        run._r.append(child)
        return run

    r(_el("w:fldChar", **{"w:fldCharType": "begin"}))
    t = _el("w:instrText", **{"xml:space": "preserve"})
    t.text = f" {instr} "
    r(t)
    r(_el("w:fldChar", **{"w:fldCharType": "separate"}))
    run = paragraph.add_run(cached)
    style_run(run, size=size, color=color, rtl=False)
    r(_el("w:fldChar", **{"w:fldCharType": "end"}))


def add_internal_link(paragraph, anchor, text, size=10.5, bold=False, color="1B2533"):
    h = _el("w:hyperlink", **{"w:anchor": anchor, "w:history": 1})
    run = paragraph.add_run(text)
    style_run(run, size=size, bold=bold, color=color)
    h.append(run._r)
    paragraph._p.append(h)


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(_el("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": fill}))


def cell_borders(cell, start=None, others="nil"):
    tcPr = cell._tc.get_or_add_tcPr()
    b = _el("w:tcBorders")
    for side in ("top", "bottom", "end", "start"):
        if side == "start" and start:
            b.append(_el(f"w:{side}", **{"w:val": "single", "w:sz": 24, "w:color": start}))
        else:
            b.append(_el(f"w:{side}", **{"w:val": others}))
    tcPr.append(b)


def table_bidi(table, width_pct=100):
    tblPr = table._tbl.tblPr
    tblPr.append(_el("w:bidiVisual"))
    tblPr.append(_el("w:tblW", **{"w:w": width_pct * 50, "w:type": "pct"}))


def png_size(path):
    with open(path, "rb") as f:
        head = f.read(24)
    return struct.unpack(">II", head[16:24])


# ───────────── document parts ─────────────
def setup_document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.header_distance = sec.footer_distance = Cm(0.9)
    sec._sectPr.append(_el("w:bidi"))
    sec.different_first_page_header_footer = True
    # Base styles: Arial for Latin and complex scripts, Arabic proofing.
    for name in ("Normal", "Heading 1", "Heading 2", "Heading 3", "Title"):
        st = doc.styles[name]
        st.font.name = FONT
        rPr = st.element.get_or_add_rPr()
        fonts = rPr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = _el("w:rFonts")
            rPr.append(fonts)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(a), FONT)
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if fonts.get(qn(a)) is not None:
                del fonts.attrib[qn(a)]
        rPr.append(_el("w:lang", **{"w:val": "en-US", "w:bidi": "ar-EG"}))
        pPr = st.element.get_or_add_pPr()
        pPr.append(_el("w:bidi"))
    doc.styles["Normal"].font.size = Pt(10.5)
    for name, size, color in (("Heading 1", 20, NAVY), ("Heading 2", 14, NAVY), ("Heading 3", 12, "24456F")):
        st = doc.styles[name]
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.element.get_or_add_rPr().append(_el("w:szCs", **{"w:val": size * 2}))
        st.element.get_or_add_rPr().append(_el("w:bCs"))
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.space_before = Pt(14 if name != "Heading 1" else 0)
        st.paragraph_format.space_after = Pt(6)
    # Compatibility: let Word lay out Arabic properly.
    settings = doc.settings.element
    settings.append(_el("w:themeFontLang", **{"w:val": "en-US", "w:bidi": "ar-EG"}))
    return doc


def header_footer(doc):
    sec = doc.sections[0]
    # Running header: logo icon + title + release.
    hp = sec.header.paragraphs[0]
    set_bidi(hp)
    run = hp.add_run()
    run.add_picture(str(ASSETS / "oms-icon.png"), height=Cm(0.55))
    r = hp.add_run("  دليل مستخدم OMS")
    style_run(r, size=8.5, bold=True, color=NAVY)
    r = hp.add_run(f"   ·   الإصدار {META['version']}   ·   {META['date']}")
    style_run(r, size=8.5, color=GREY)
    pPr = hp._p.get_or_add_pPr()
    bdr = _el("w:pBdr")
    bdr.append(_el("w:bottom", **{"w:val": "single", "w:sz": 6, "w:space": 4, "w:color": BLUE}))
    pPr.append(bdr)
    # Footer: page X of Y + SHA.
    fp = sec.footer.paragraphs[0]
    set_bidi(fp, center=True)
    r = fp.add_run("صفحة ")
    style_run(r, size=8.5, color=GREY)
    add_field(fp, "PAGE", "1", size=8.5)
    r = fp.add_run(" من ")
    style_run(r, size=8.5, color=GREY)
    add_field(fp, "NUMPAGES", "1", size=8.5)
    r = fp.add_run(f"     ·     SHA {META['sha']}")
    style_run(r, size=7.5, color="8A97AA", rtl=False)


def cover(doc):
    p = para(doc, center=True, space_after=0)
    p.paragraph_format.space_before = Pt(70)
    p.add_run().add_picture(str(ASSETS / "oms-logo.png"), width=Cm(13))
    # Coloured stripe (a one-cell table with a navy fill).
    t = doc.add_table(rows=1, cols=1)
    table_bidi(t)
    c = t.cell(0, 0)
    shade_cell(c, NAVY)
    cell_borders(c)
    cp = c.paragraphs[0]
    set_bidi(cp)
    cp.paragraph_format.space_before = Pt(18)
    add_runs(cp, [{"t": "دليل المستخدم", "b": True}], size=30, color="FFFFFF")
    sub = para(c, [{"t": "نظام إدارة العمليات OMS — شرح الشاشات والمهام حسب الدور الوظيفي", "b": False}], size=13, color="BFD3F2", space_after=14)
    for k, v in (("الإصدار", META["version"]), ("رقم البناء (Git SHA)", META["sha"]), ("تاريخ الإعداد", META["date"]), ("اللغة", "العربية")):
        mp = para(c, space_after=3)
        add_runs(mp, [{"t": f"{k}:  "}], size=11, color="9FB6D8")
        add_runs(mp, [{"t": v, "code": False, "b": True}], size=11, color="FFFFFF")
    last = para(c, space_after=18)
    del sub, last
    note = para(doc, [{"t": "صور هذا الدليل مأخوذة من النظام الفعلي ببيانات تجريبية مُختلقة بالكامل (شركة الأفق للتجارة). لا تحتوي الصور على كلمات مرور أو بيانات عملاء حقيقيين."}], size=9, color=GREY, space_after=0)
    note.paragraph_format.space_before = Pt(40)
    note.add_run().add_break(WD_BREAK.PAGE)


def toc(doc, manual):
    h = doc.add_paragraph(style="Heading 1")
    set_bidi(h)
    add_runs(h, [{"t": "المحتويات"}], size=20, color=NAVY, bold=True)
    add_bookmark(h, "toc")
    count = 0
    for ch in manual["chapters"]:
        h1 = ch["blocks"][0]
        p = para(doc, space_after=2)
        p.paragraph_format.space_before = Pt(5)
        add_internal_link(p, h1["id"], f"{ch['no']}. {ch['title']}", size=11, bold=True, color=NAVY)
        r = p.add_run("   ·   ص ")
        style_run(r, size=9, color=GREY)
        add_field(p, f"PAGEREF {h1['id']} \\h", "", size=9, color=BLUE)
        count += 1
        for b in ch["blocks"]:
            if b["type"] == "h2":
                p = para(doc, space_after=1)
                indent_start(p, 420)
                add_internal_link(p, b["id"], f"{b['no']} {b['text']}", size=9.5, color="3A4658")
                r = p.add_run("   ·   ص ")
                style_run(r, size=8.5, color=GREY)
                add_field(p, f"PAGEREF {b['id']} \\h", "", size=8.5, color=BLUE)
                count += 1
    hint = para(doc, [{"t": "انقر أي عنوان للانتقال إليه. أرقام الصفحات في هذا الفهرس تتحدث في Word بالضغط على Ctrl+A ثم F9."}], size=8.5, color=GREY)
    hint.paragraph_format.space_before = Pt(8)
    return count


def figure(doc, b):
    path = SHOTS / f"{b['shot']}.png"
    w, h = png_size(path)
    width_cm = figure_width_mm(b["shot"], w, h, max_h=195.0) / 10
    p = para(doc, center=True, space_after=2, keep_next=True)
    p.paragraph_format.space_before = Pt(6)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    cap = para(doc, [{"t": f"شكل {b['no']}: {b['caption']}", "b": True}], size=9, center=True, space_after=6, color="3A4658")
    cap.paragraph_format.keep_with_next = True  # keep the caption with its legend


def box(doc, b):
    t = doc.add_table(rows=1, cols=1)
    table_bidi(t)
    c = t.cell(0, 0)
    shade_cell(c, BOX_FILL.get(b["kind"], "EEF4FD"))
    cell_borders(c, start=BOX_LINE.get(b["kind"], BLUE))
    tp = c.paragraphs[0]
    set_bidi(tp)
    tp.paragraph_format.keep_with_next = True
    add_runs(tp, [{"t": b["title"], "b": True}], size=10)
    for x in b["body"]:
        if x["type"] == "li":
            p = para(c, [{"t": "•  "}] + x["runs"], size=9.8, space_after=2)
            indent_start(p, 240, 240)
        else:
            para(c, x["runs"], size=9.8, space_after=2)
    para(doc, space_after=2)


def data_table(doc, b):
    cols = len(b["header"])
    t = doc.add_table(rows=1 + len(b["rows"]), cols=cols)
    t.style = doc.styles["Table Grid"]
    table_bidi(t)
    # Repeat the header row on every page and keep rows whole.
    trPr = t.rows[0]._tr.get_or_add_trPr()
    trPr.append(_el("w:tblHeader"))
    for j, cell_runs in enumerate(b["header"]):
        c = t.cell(0, j)
        shade_cell(c, NAVY)
        cp = c.paragraphs[0]
        set_bidi(cp)
        add_runs(cp, cell_runs, size=9.5, color="FFFFFF", bold=True)
    for i, row in enumerate(b["rows"], start=1):
        t.rows[i]._tr.get_or_add_trPr().append(_el("w:cantSplit"))
        for j in range(cols):
            c = t.cell(i, j)
            if i % 2 == 0:
                shade_cell(c, "F5F8FC")
            cp = c.paragraphs[0]
            set_bidi(cp)
            cp.paragraph_format.space_after = Pt(1)
            add_runs(cp, row[j] if j < len(row) else [], size=9.5)
    para(doc, space_after=2)


def legend(doc, b):
    for it in b["items"]:
        p = para(doc, space_after=2)
        r = p.add_run(f" {it['n']} ")
        style_run(r, size=8.5, bold=True, color="FFFFFF", rtl=False)
        r._r.get_or_add_rPr().append(_el("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": ORANGE}))
        r = p.add_run("  ")
        style_run(r, size=9)
        add_runs(p, it["runs"], size=10)
        indent_start(p, 360, 360)


def steps(doc, b):
    for it in b["items"]:
        p = para(doc, space_after=2)
        r = p.add_run(f"{it['n']}. ")
        style_run(r, size=10.5, bold=True, color=BLUE)
        add_runs(p, it["runs"])
        indent_start(p, 360, 360)
        p.paragraph_format.keep_together = True
        for s in it["sub"]:
            sp = para(doc, [{"t": "–  "}] + s, size=10, space_after=1)
            indent_start(sp, 720, 240)


def bullets(doc, b):
    for it in b["items"]:
        p = para(doc, [{"t": "•  "}] + it["runs"], space_after=2)
        indent_start(p, 360, 240)


def build_docx(manual, out_path):
    doc = setup_document()
    header_footer(doc)
    cover(doc)
    toc_count = toc(doc, manual)
    headings = 0
    images = 0
    for ch in manual["chapters"]:
        for b in ch["blocks"]:
            t = b["type"]
            if t == "h1":
                h = doc.add_paragraph(style="Heading 1")
                set_bidi(h)
                h.paragraph_format.page_break_before = True
                add_runs(h, [{"t": f"{ch['no']}. {b['text']}"}], size=20, color=NAVY, bold=True)
                add_bookmark(h, b["id"])
                pPr = h._p.get_or_add_pPr()
                bdr = _el("w:pBdr")
                bdr.append(_el("w:bottom", **{"w:val": "single", "w:sz": 18, "w:space": 4, "w:color": BLUE}))
                pPr.append(bdr)
                headings += 1
            elif t == "h2":
                h = doc.add_paragraph(style="Heading 2")
                set_bidi(h)
                add_runs(h, [{"t": f"{b['no']} {b['text']}"}], size=14, color=NAVY, bold=True)
                add_bookmark(h, b["id"])
                headings += 1
            elif t == "h3":
                h = doc.add_paragraph(style="Heading 3")
                set_bidi(h)
                add_runs(h, [{"t": b["text"]}], size=12, color="24456F", bold=True)
                add_bookmark(h, b["id"])
            elif t == "p":
                para(doc, b["runs"])
            elif t == "figure":
                figure(doc, b)
                images += 1
            elif t == "legend":
                legend(doc, b)
            elif t == "ol":
                steps(doc, b)
            elif t == "ul":
                bullets(doc, b)
            elif t == "table":
                data_table(doc, b)
            elif t == "box":
                box(doc, b)
    # Ask Word to refresh fields (PAGEREF / NUMPAGES) when the file is opened.
    doc.settings.element.append(_el("w:updateFields", **{"w:val": "true"}))
    doc.core_properties.title = f"{META['title']} — {META['version']}"
    doc.core_properties.subject = f"OMS {META['version']} ({META['sha']})"
    doc.core_properties.language = "ar-EG"
    doc.save(out_path)
    return {"toc_entries": toc_count, "headings_bookmarked": headings, "images": images}
